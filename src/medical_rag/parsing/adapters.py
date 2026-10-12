from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from pathlib import PurePath
from typing import BinaryIO, Mapping

import xlrd

from medical_rag.domain.enums import EvidenceKind
from medical_rag.parsing.contracts import (
    OCRParser,
    ParsedBlock,
    ParsedDocument,
    ParsedSegment,
    ParsedTable,
    TextParser,
    as_bytes_io,
)
from medical_rag.parsing.filetypes import IMAGE_FILE_TYPES

# 解析适配器专属的页级块 ID 前缀（单一来源；不要散落 f-string）。
# 与 LayoutNormalizer/TextSourceParser 的 make_block_id（blk-NNNN）分工：
# 块 ID 只在解析适配器用于 disambiguate，同一文档内保证唯一（页内从 0 起）。
PAGE_BLOCK_ID_FORMAT = "p{page}-b{index}"

# 各适配器产出版本（单一来源，服务端透传）。只提升本任务实际改动的解析器：
# PDF / DOCX / 电子表格 / OCR 均改为 v2 版面块输出。
PARSER_VERSION_PDF = "pdf-parser-v2"
PARSER_VERSION_DOCX = "docx-parser-v2"
PARSER_VERSION_SPREADSHEET = "spreadsheet-parser-v2"
PARSER_VERSION_OCR = "ocr-parser-v2"

# OCR 回退策略：单页无原生文本或有效字符数低于该阈值时判定需要 OCR。
_OCR_PAGE_MIN_CHARS = 10

# OCR 后端的坐标型 source_anchor（x:y:w:h），解析为 bbox（fail-closed：不匹配则无 bbox）。
_BBOX_ANCHOR_PATTERN = re.compile(
    r"^x(?P<x>-?\d+(?:\.\d+)?):y(?P<y>-?\d+(?:\.\d+)?):"
    r"w(?P<w>-?\d+(?:\.\d+)?):h(?P<h>-?\d+(?:\.\d+)?)$"
)


def _page_block_id(page: int, index: int) -> str:
    return PAGE_BLOCK_ID_FORMAT.format(page=page, index=index)


def _bbox_from_anchor(source_anchor: str | None) -> tuple[float, float, float, float] | None:
    if not source_anchor:
        return None
    match = _BBOX_ANCHOR_PATTERN.match(source_anchor)
    if not match:
        return None
    return tuple(float(match.group(name)) for name in ("x", "y", "w", "h"))  # type: ignore[return-value]


def _heading_level(paragraph: object) -> int | None:
    """分类 Word 标题样式并返回层级（1 起）；非标题返回 None。

    兼容英文 'Heading N' 与中文 '标题 N' 内建样式名。
    """
    style = getattr(paragraph, "style", None)
    for current in (style, getattr(style, "base_style", None)):
        name = (getattr(current, "name", None) or "").strip()
        match = re.match(r"^(?:Heading|标题)\s*([1-9])$", name, re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None


def detect_language(text: str) -> str:
    has_chinese = any("\u4e00" <= char <= "\u9fff" for char in text)
    has_latin = any(char.isascii() and char.isalpha() for char in text)
    if has_chinese and has_latin:
        return "mixed"
    if has_chinese:
        return "zh-CN"
    if has_latin:
        return "en-US"
    return "und"


def _segment_from_mapping(item: Mapping[str, object], default_kind: EvidenceKind) -> ParsedSegment:
    text = str(item.get("text", ""))
    return ParsedSegment(
        text=text,
        page_number=int(item["page_number"]) if item.get("page_number") is not None else None,
        sheet_name=str(item["sheet_name"]) if item.get("sheet_name") is not None else None,
        cell_range=str(item["cell_range"]) if item.get("cell_range") is not None else None,
        source_anchor=(
            str(item["source_anchor"]) if item.get("source_anchor") is not None else None
        ),
        language=str(item.get("language") or detect_language(text)),
        confidence=float(item.get("confidence", 1.0)),
        kind=EvidenceKind(item.get("kind", default_kind)),
    )


def _build_ocr_blocks(document: ParsedDocument) -> tuple[ParsedBlock, ...]:
    """从 OCR 段派生版面块：坐标 source_anchor → bbox、confidence → confidence。

    单页图像记 page_number=1（OCR 输入视为单页）；无坐标/缺 confidence 时
    fail-closed（bbox=None, confidence 沿用段值）。不在此引入 ocr_noise 语义
    （那是 DocumentCleaner 的词汇）。
    """
    blocks: list[ParsedBlock] = []
    for index, segment in enumerate(document.segments):
        blocks.append(
            ParsedBlock(
                block_id=_page_block_id(1, index),
                block_type="paragraph",
                raw_text=segment.text,
                page_number=1,
                bbox=_bbox_from_anchor(segment.source_anchor),
                reading_order=index,
                language=segment.language,
                confidence=segment.confidence,
                source_anchor=segment.source_anchor,
            )
        )
    return tuple(blocks)


def render_xls_cell_text(cell_type: int, value: object, datemode: int) -> str:
    """按 xlrd ``ctype`` 还原一个 ``.xls`` 单元格的文本。

    ``.xls`` 与 ``.xlsx`` 的差别在于类型信息只能从 xlrd 的 ``ctype`` 拿到：
    ``XL_CELL_DATE`` 的 value 是 Excel 日期序列号（float），必须经
    :func:`xlrd.xldate_as_datetime` 还原成真实日期，否则证据里会出现
    ``45000.0`` 这种"日期"。数字按整数值省略小数点（``105.0`` → ``105``），
    布尔按 Excel 显示习惯输出 ``TRUE``/``FALSE``。
    """
    if value is None:
        return ""
    if cell_type == xlrd.XL_CELL_DATE:
        try:
            moment = xlrd.xldate_as_datetime(float(value), datemode)
        except (ValueError, TypeError):
            return str(value)
        # 纯日期（时间部分为 0）只输出日期，避免证据里出现无意义的 00:00:00。
        return moment.date().isoformat() if moment.time() == time(0, 0) else moment.isoformat(sep=" ")
    if cell_type == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if float(value) else "FALSE"
    if cell_type in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if cell_type == xlrd.XL_CELL_NUMBER:
        number = float(value)
        return str(int(number)) if number.is_integer() else repr(number)
    return str(value)


def should_ocr_page(blocks: tuple[ParsedBlock, ...], page_number: int) -> bool:
    """单页 OCR 回退判定：仅当该页无可用原生文本或覆盖度低于配置阈值时返回 True。

    逐页判定 => 不会因为只有一页空白就对整份 PDF 触发 OCR。页内有效文本字符数
    （去掉空白）低于 _OCR_PAGE_MIN_CHARS 视作覆盖不足。

    **表格块的原生文本同样计入页面有效字符数**：整页原生表格（检验单）是完全可用
    的原生内容，若把它记为 0 字符，混合 PDF 会被误判为"缺文本"而整页重扫，OCR 结果
    与原生表格会以不同锚点（``p1`` vs ``p1:text``）双份进入索引。
    """
    page_blocks = [block for block in blocks if block.page_number == page_number]
    if not page_blocks:
        return True
    total_chars = sum(
        len("".join(block.raw_text.split())) for block in page_blocks
    )
    return total_chars < _OCR_PAGE_MIN_CHARS


@dataclass(frozen=True, slots=True)
class ParserAdapters:
    text: TextParser | None = None
    ocr: OCRParser | None = None
    table: object | None = None
    # Task 11 seam: per-page OCR backend.  When present, mixed PDFs are backfilled
    # page-by-page instead of raising MixedPdfOcrDeferredError (see
    # medical_rag.parsing.per_page_ocr.ocr_backfill_pdf).  Additive and optional.
    per_page_ocr: object | None = None
    # Remote per-PDF OCR backend *factory* (from build_remote_ocr).  When present
    # it must be a callable ``factory(pdf_bytes) -> RemotePerPageOCRBackend`` and
    # is built lazily inside the service per document.  Additive and optional; the
    # deterministic ``per_page_ocr`` fixture backend is untouched.
    remote_ocr: object | None = None


def build_parser_adapters(settings: object) -> ParserAdapters:
    """Composition glue: ``Settings`` → :class:`ParserAdapters` for ParsingService.

    When ``ocr_service_url`` is set, wires the remote PaddleOCR-VL backend: the
    whole-doc :class:`RemoteOCRParser` as ``ocr`` (image path) plus the per-PDF
    backend *factory* as ``remote_ocr`` (scanned-PDF backfill).  When the URL is
    empty (the default), returns an empty adapter set so ``ParsingService`` uses
    only its deterministic native parsers and the remote backend stays disabled.

    This is deliberately a lazy import of ``remote_ocr`` so wireing remote OCR
    never forces ``pypdfium2``/``httpx`` side effects at module import time.
    """
    from medical_rag.parsing.remote_ocr import build_remote_ocr

    parser, backend_factory = build_remote_ocr(settings)
    if parser is None or backend_factory is None:
        return ParserAdapters()
    return ParserAdapters(ocr=parser, remote_ocr=backend_factory)


class DeterministicTextParser:
    def __init__(self, fixtures: Mapping[str, list[Mapping[str, object]]]) -> None:
        self.fixtures = fixtures

    def parse(self, file: BinaryIO, file_type: str) -> ParsedDocument:
        del file
        extension = PurePath(file_type).suffix.lower()
        if not extension:
            extension = {
                "application/pdf": ".pdf",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
            }.get(file_type, file_type)
        items = self.fixtures.get(file_type, self.fixtures.get(extension, []))
        segments = tuple(_segment_from_mapping(item, EvidenceKind.REPORT_TEXT) for item in items)
        raw_text = "\n".join(segment.text for segment in segments)
        return ParsedDocument(
            file_type=file_type,
            raw_text=raw_text,
            segments=segments,
            language=detect_language(raw_text),
        )


class DeterministicOCRParser:
    def __init__(self, fixtures: Mapping[str, list[Mapping[str, object]]]) -> None:
        self.fixtures = fixtures

    def parse_images(self, file: BinaryIO, language_hints: list[str]) -> ParsedDocument:
        del file
        items = self.fixtures.get("image", [])
        segments = tuple(
            _segment_from_mapping({"confidence": 0.95, **item}, EvidenceKind.REPORT_TEXT)
            for item in items
        )
        raw_text = "\n".join(segment.text for segment in segments)
        language = detect_language(raw_text)
        if language == "und" and language_hints:
            language = language_hints[0]
        document = ParsedDocument(
            file_type="image",
            raw_text=raw_text,
            segments=segments,
            language=language,
        )
        return replace(
            document,
            parser_version=PARSER_VERSION_OCR,
            blocks=_build_ocr_blocks(document),
        )


class PdfTextParser:
    supported_file_types = frozenset({".pdf", "application/pdf"})

    def parse(self, file: BinaryIO, file_type: str) -> ParsedDocument:
        from pypdf import PdfReader

        content = as_bytes_io(file)
        reader = PdfReader(content)
        segments: list[ParsedSegment] = []
        blocks: list[ParsedBlock] = []
        block_order = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            segments.append(
                ParsedSegment(
                    text=text,
                    page_number=page_number,
                    source_anchor=f"p{page_number}:text",
                    language=detect_language(text),
                )
            )
            # 每页一个版面块；块 ID 页级唯一（p1-b0, p2-b0…）。source_anchor
            # 只在确实存在时透传（fail-closed）。PDF 无标题结构 => 空 section_path。
            blocks.append(
                ParsedBlock(
                    block_id=_page_block_id(page_number, 0),
                    block_type="paragraph",
                    raw_text=text,
                    page_number=page_number,
                    section_path=(),
                    reading_order=block_order,
                    language=detect_language(text),
                    confidence=1.0,
                    source_anchor=f"p{page_number}:text",
                )
            )
            block_order += 1
        raw_text = "\n".join(segment.text for segment in segments)
        return ParsedDocument(
            file_type=file_type,
            raw_text=raw_text,
            segments=tuple(segments),
            blocks=tuple(blocks),
            language=detect_language(raw_text),
            parser_version=PARSER_VERSION_PDF,
        )


class WordTextParser:
    supported_file_types = frozenset(
        {
            ".docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }
    )

    def parse(self, file: BinaryIO, file_type: str) -> ParsedDocument:
        from docx import Document as WordDocument

        document = WordDocument(as_bytes_io(file))
        segments: list[ParsedSegment] = []
        tables: list[ParsedTable] = []
        blocks: list[ParsedBlock] = []
        # 标题样式 → 章节路径；标题只分轴、段落/表格承接当前栈。
        section_stack: list[str] = []
        paragraph_index = 0
        table_index = 0
        reading_order = 0

        # 按正文出现顺序遍历（段落与表格交错），保证 reading_order 反映版面顺序。
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                from docx.text.paragraph import Paragraph

                paragraph = Paragraph(child, document)
                paragraph_index += 1
                text = paragraph.text.strip()
                if not text:
                    continue
                level = _heading_level(paragraph)
                if level is not None:
                    # 标题块：更新章节栈并承载自身标题文本。
                    heading_text = text
                    section_stack = section_stack[: level - 1] + [heading_text]
                    segments.append(
                        ParsedSegment(
                            text=text,
                            source_anchor=f"doc:paragraph:{paragraph_index}",
                            language=detect_language(text),
                        )
                    )
                    blocks.append(
                        ParsedBlock(
                            block_id=_page_block_id(1, reading_order),
                            block_type="heading",
                            raw_text=text,
                            section_path=tuple(section_stack),
                            reading_order=reading_order,
                            language=detect_language(text),
                            confidence=1.0,
                            source_anchor=f"doc:paragraph:{paragraph_index}",
                        )
                    )
                    reading_order += 1
                    continue
                segments.append(
                    ParsedSegment(
                        text=text,
                        source_anchor=f"doc:paragraph:{paragraph_index}",
                        language=detect_language(text),
                    )
                )
                blocks.append(
                    ParsedBlock(
                        block_id=_page_block_id(1, reading_order),
                        block_type="paragraph",
                        raw_text=text,
                        section_path=tuple(section_stack),
                        reading_order=reading_order,
                        language=detect_language(text),
                        confidence=1.0,
                        source_anchor=f"doc:paragraph:{paragraph_index}",
                    )
                )
                reading_order += 1
            elif child.tag.endswith("}tbl"):
                from docx.table import Table

                table = Table(child, document)
                table_index += 1
                rows = tuple(tuple(cell.text.strip() for cell in row.cells) for row in table.rows)
                if not rows:
                    continue
                parsed_table = ParsedTable(
                    headers=rows[0],
                    rows=rows[1:],
                    source_anchor=f"doc:table:{table_index}",
                    language=detect_language("\n".join(" ".join(row) for row in rows)),
                )
                tables.append(parsed_table)
                segments.append(
                    ParsedSegment(
                        text=parsed_table.text,
                        source_anchor=parsed_table.source_anchor,
                        language=parsed_table.language,
                        kind=EvidenceKind.REPORT_TABLE,
                    )
                )
                blocks.append(
                    ParsedBlock(
                        block_id=_page_block_id(1, reading_order),
                        block_type="table",
                        raw_text=parsed_table.text,
                        section_path=tuple(section_stack),
                        reading_order=reading_order,
                        language=parsed_table.language,
                        confidence=1.0,
                        source_anchor=parsed_table.source_anchor,
                    )
                )
                reading_order += 1
        raw_text = "\n".join(segment.text for segment in segments)
        return ParsedDocument(
            file_type=file_type,
            raw_text=raw_text,
            segments=tuple(segments),
            tables=tuple(tables),
            blocks=tuple(blocks),
            language=detect_language(raw_text),
            parser_version=PARSER_VERSION_DOCX,
        )


# ---------------------------------------------------------------------------
# 电子表格（检验单）单元格渲染 + 真实 used region / 表头行探测
# ---------------------------------------------------------------------------

# 表头行探测只看前 N 行（标题行通常在前几行内）。
_SPREADSHEET_HEADER_SCAN_ROWS = 5

_HAS_DATE_TOKEN = re.compile(r"(?i)[ymdhs]")
_NUMERIC_PLACEHOLDER = re.compile(r"[#0?][#0?,]*(?:\.[#0?]+)?")


def _strip_spreadsheet_format(number_format: str) -> str:
    """取第一个 format section 并去掉条件/颜色段与引号字面量。"""
    section = (number_format or "").split(";")[0]
    section = re.sub(r"\[[^\]]*\]", "", section)
    section = re.sub(r'"([^"]*)"', r"\1", section)
    return section.replace("\\", "")


def _render_excel_datetime(value: datetime | date | time, pattern: str) -> str:
    """按 Excel 日期格式串渲染（支持 y/m/d/h/s 与 AM/PM 的实际使用子集）。

    只实现检验单会遇到的格式（``yyyy/m/d``、``yyyy-mm-dd hh:mm:ss`` …）：
    未识别的字母按字面量输出，绝不臆造格式。
    """
    hour = getattr(value, "hour", 0)
    minute = getattr(value, "minute", 0)
    second = getattr(value, "second", 0)
    rendered: list[str] = []
    index = 0
    seen_hour = False
    while index < len(pattern):
        char = pattern[index]
        lower = char.lower()
        if not lower.isalpha():
            rendered.append(char)
            index += 1
            continue
        run = 1
        while (
            index + run < len(pattern)
            and pattern[index + run].lower() == lower
        ):
            run += 1
        if lower == "y":
            rendered.append(f"{value.year:04d}" if run >= 3 else f"{value.year % 100:02d}")
        elif lower == "m":
            if seen_hour:
                rendered.append(f"{minute:02d}" if run >= 2 else str(minute))
            else:
                rendered.append(f"{value.month:02d}" if run >= 2 else str(value.month))
        elif lower == "d":
            rendered.append(f"{value.day:02d}" if run >= 2 else str(value.day))
        elif lower == "h":
            seen_hour = True
            rendered.append(f"{hour:02d}" if run >= 2 else str(hour))
        elif lower == "s":
            rendered.append(f"{second:02d}" if run >= 2 else str(second))
        else:
            rendered.append(char * run)
        index += run
    return "".join(rendered)


def _render_excel_number(value: int | float, pattern: str) -> str | None:
    """按 ``#,##0.00`` / ``0.00%`` 这类数字格式渲染；无法识别返回 None。"""
    placeholder = _NUMERIC_PLACEHOLDER.search(pattern)
    if placeholder is None:
        return None
    body = placeholder.group(0)
    integer_part, _, decimals = body.partition(".")
    grouping = "," in integer_part
    scaled = float(value) * 100 if "%" in pattern else float(value)
    text = f"{scaled:,.{len(decimals)}f}" if grouping else f"{scaled:.{len(decimals)}f}"
    return f"{pattern[: placeholder.start()]}{text}{pattern[placeholder.end() :]}"


def _format_spreadsheet_cell(value: object, number_format: str) -> str:
    """把 openpyxl 的单元格值渲染为**用户在 Excel 里看到的**显示文本。

    键入格式（日期/千分位/百分比/货币）优先；无法识别格式时回退到值的字符形式，
    绝不为了"好看"丢掉数值本身。
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    pattern = _strip_spreadsheet_format(number_format)
    if isinstance(value, (datetime, date, time)):
        if _HAS_DATE_TOKEN.search(pattern):
            return _render_excel_datetime(value, pattern)
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, (int, float)):
        if pattern and pattern.lower() != "general":
            rendered = _render_excel_number(value, pattern)
            if rendered is not None:
                return rendered
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


@dataclass(frozen=True, slots=True)
class _SheetGrid:
    """一张工作表的显示文本网格（绝对坐标锚定，合并单元格已展开）。

    ``raw_rows`` 保留**未展开**视图（合并区只有左上角有值），用于表头行探测与标题
    文本：跨列合并的标题行在 raw 视图里只有 1 个非空单元格，因此不会被误判为
    "非空列数最多"的表头行。
    """

    sheet_name: str
    expanded_rows: tuple[tuple[str, ...], ...]
    raw_rows: tuple[tuple[str, ...], ...]
    first_row: int
    first_col: int
    header_index: int

    @property
    def header_row(self) -> int:
        return self.first_row + self.header_index

    @property
    def last_row(self) -> int:
        return self.first_row + len(self.expanded_rows) - 1

    @property
    def last_col(self) -> int:
        return self.first_col + len(self.expanded_rows[0]) - 1

    @property
    def table_cell_range(self) -> str:
        return self._range(self.header_row, self.last_row, self.first_col, self.last_col)

    def range_for_relative_row(self, index: int) -> str:
        """某一相对行（0 起）的显示区间；合并展开后的宽度决定区间右端。"""
        last_col = max(
            (
                self.first_col + offset
                for offset, cell in enumerate(self.expanded_rows[index])
                if cell.strip()
            ),
            default=self.first_col,
        )
        absolute_row = self.first_row + index
        return self._range(absolute_row, absolute_row, self.first_col, last_col)

    @staticmethod
    def _range(start_row: int, end_row: int, start_col: int, end_col: int) -> str:
        from openpyxl.utils import get_column_letter

        return (
            f"{get_column_letter(start_col)}{start_row}:"
            f"{get_column_letter(end_col)}{end_row}"
        )


def _grid_from_cells(
    sheet_name: str,
    cells: dict[tuple[int, int], str],
    spans: list[tuple[int, int, int, int]],
) -> _SheetGrid | None:
    """由 ``{(row, col): text}``（1 起绝对坐标）+ 合并跨度构造显示网格。"""
    if not cells and not spans:
        return None
    rows = [row for row, _ in cells] + [span[0] for span in spans]
    cols = [col for _, col in cells] + [span[1] for span in spans]
    end_rows = [row for row, _ in cells] + [span[2] for span in spans]
    end_cols = [col for _, col in cells] + [span[3] for span in spans]
    first_row, last_row = min(rows), max(end_rows)
    first_col, last_col = min(cols), max(end_cols)
    expanded = dict(cells)
    for top, left, bottom, right in spans:
        value = cells.get((top, left), "")
        if not value.strip():
            continue
        for row in range(top, bottom + 1):
            for col in range(left, right + 1):
                expanded.setdefault((row, col), value)

    def plan(source: dict[tuple[int, int], str]) -> tuple[tuple[str, ...], ...]:
        return tuple(
            tuple(source.get((row, col), "") for col in range(first_col, last_col + 1))
            for row in range(first_row, last_row + 1)
        )

    raw_rows = plan(cells)
    return _SheetGrid(
        sheet_name=sheet_name,
        expanded_rows=plan(expanded),
        raw_rows=raw_rows,
        first_row=first_row,
        first_col=first_col,
        header_index=_detect_header_index(raw_rows),
    )


def _detect_header_index(rows: tuple[tuple[str, ...], ...]) -> int:
    """在前 N 行里挑表头行：非空列数最多、其后各行列数最一致，平局取最早。

    标题行（如 A1:C1 合并的"血常规检验报告单"）非空列数少、后续行宽度不一致，
    因此不会被当成表头——这是"标题行不得变成每个子块的表头"的判据。
    """
    limit = min(len(rows), _SPREADSHEET_HEADER_SCAN_ROWS)
    best_index = 0
    best_key: tuple[int, float, int] | None = None
    for index in range(limit):
        width = sum(1 for cell in rows[index] if cell.strip())
        body = [sum(1 for cell in row if cell.strip()) for row in rows[index + 1 :]]
        consistency = (
            Counter(body).most_common(1)[0][1] / len(body) if body else 0.0
        )
        key = (width, consistency, -index)
        if best_key is None or key > best_key:
            best_key, best_index = key, index
    return best_index


def _split_sheet_grid(
    grid: _SheetGrid,
) -> tuple[list[str], tuple[str, ...], list[tuple[str, ...]]]:
    """拆成（标题行文本、表头行、数据行）。

    标题文本取**未展开**视图（避免跨列合并值被重复 N 次），数据行取展开视图
    （合并的指标名必须在每一行都存在）。
    """
    titles: list[str] = []
    for row in grid.raw_rows[: grid.header_index]:
        cells = list(dict.fromkeys(cell.strip() for cell in row if cell.strip()))
        if cells:
            titles.append(" ".join(cells))
    headers = grid.expanded_rows[grid.header_index]
    data_rows = [row for row in grid.expanded_rows[grid.header_index + 1 :] if any(row)]
    return titles, headers, data_rows


def _xlsx_sheet_grid(sheet: object) -> _SheetGrid | None:
    cells: dict[tuple[int, int], str] = {}
    for row in sheet.iter_rows():  # type: ignore[attr-defined]
        for cell in row:
            text = _format_spreadsheet_cell(cell.value, cell.number_format or "")
            if text.strip():
                cells[(cell.row, cell.column)] = text
    spans = [
        (merged.min_row, merged.min_col, merged.max_row, merged.max_col)
        for merged in sheet.merged_cells.ranges  # type: ignore[attr-defined]
    ]
    return _grid_from_cells(sheet.title, cells, spans)  # type: ignore[attr-defined]


def _xls_sheet_grid(sheet: object, datemode: int) -> _SheetGrid | None:
    cells: dict[tuple[int, int], str] = {}
    for row in range(sheet.nrows):  # type: ignore[attr-defined]
        for col in range(sheet.ncols):  # type: ignore[attr-defined]
            cell = sheet.cell(row, col)  # type: ignore[attr-defined]
            text = render_xls_cell_text(cell.ctype, cell.value, datemode)
            if text.strip():
                cells[(row + 1, col + 1)] = text
    spans: list[tuple[int, int, int, int]] = []
    for merged in getattr(sheet, "merged_cells", ()) or ():
        # xlrd 的 merged_cells 是 (rlo, rhi, clo, chi)，且 rhi/chi 为开区间。
        rlo, rhi, clo, chi = merged
        spans.append((rlo + 1, clo + 1, rhi, chi))
    return _grid_from_cells(sheet.name, cells, spans)  # type: ignore[attr-defined]


class ExcelParser:
    supported_file_types = frozenset(
        {
            ".xlsx",
            ".xls",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        }
    )

    def parse(self, file: BinaryIO, file_type: str) -> ParsedDocument:
        normalized_type = file_type.lower().strip()
        if normalized_type in {".xls", "application/vnd.ms-excel"}:
            workbook = xlrd.open_workbook(
                file_contents=as_bytes_io(file).read(), on_demand=True
            )
            grids = [
                grid
                for grid in (
                    _xls_sheet_grid(sheet, workbook.datemode) for sheet in workbook.sheets()
                )
                if grid is not None
            ]
        else:
            from openpyxl import load_workbook

            # read_only=False + data_only=True: data_only keeps the cached formula
            # results while the non-read-only book is what exposes
            # ``cell.number_format`` and ``sheet.merged_cells`` (the P1.2 gaps).
            workbook = load_workbook(as_bytes_io(file), data_only=True)
            grids = [
                grid
                for grid in (_xlsx_sheet_grid(sheet) for sheet in workbook.worksheets)
                if grid is not None
            ]

        segments: list[ParsedSegment] = []
        tables: list[ParsedTable] = []
        blocks: list[ParsedBlock] = []
        block_order = 0
        for sheet_ordinal, grid in enumerate(grids, start=1):
            titles, headers, data_rows = _split_sheet_grid(grid)
            section_path = (titles[0],) if titles else ()
            for title_index, title in enumerate(titles):
                title_range = grid.range_for_relative_row(title_index)
                title_text = title
                language = detect_language(title_text)
                segments.append(
                    ParsedSegment(
                        text=title_text,
                        sheet_name=grid.sheet_name,
                        cell_range=title_range,
                        source_anchor=f"sheet:{grid.sheet_name}!{title_range}",
                        language=language,
                        kind=EvidenceKind.REPORT_TEXT,
                    )
                )
                blocks.append(
                    ParsedBlock(
                        block_id=_page_block_id(sheet_ordinal, len(blocks)),
                        block_type="heading",
                        raw_text=title_text,
                        sheet_name=grid.sheet_name,
                        cell_range=title_range,
                        section_path=section_path,
                        reading_order=block_order,
                        language=language,
                        confidence=1.0,
                        source_anchor=f"sheet:{grid.sheet_name}!{title_range}",
                    )
                )
                block_order += 1
            cell_range = grid.table_cell_range
            parsed_table = ParsedTable(
                headers=headers,
                rows=tuple(data_rows),
                sheet_name=grid.sheet_name,
                cell_range=cell_range,
                source_anchor=f"sheet:{grid.sheet_name}!{cell_range}",
                language=detect_language("\n".join(" ".join(row) for row in (headers, *data_rows))),
            )
            tables.append(parsed_table)
            segments.append(
                ParsedSegment(
                    text=parsed_table.text,
                    sheet_name=grid.sheet_name,
                    cell_range=cell_range,
                    source_anchor=parsed_table.source_anchor,
                    language=parsed_table.language,
                    kind=EvidenceKind.REPORT_TABLE,
                )
            )
            # 电子表格块按 sheet 页级唯一；block_type=table，保留 sheet_name 与**真实**
            # cell_range（表头行 + 数据行，絶不从 A1 强推）；标题行单独成为 heading 块。
            blocks.append(
                ParsedBlock(
                    block_id=_page_block_id(sheet_ordinal, len(blocks)),
                    block_type="table",
                    raw_text=parsed_table.text,
                    sheet_name=grid.sheet_name,
                    cell_range=cell_range,
                    section_path=section_path,
                    reading_order=block_order,
                    language=parsed_table.language,
                    confidence=1.0,
                    source_anchor=parsed_table.source_anchor,
                )
            )
            block_order += 1
        raw_text = "\n".join(segment.text for segment in segments)
        return ParsedDocument(
            file_type=file_type,
            raw_text=raw_text,
            segments=tuple(segments),
            tables=tuple(tables),
            blocks=tuple(blocks),
            language=detect_language(raw_text),
            parser_version=PARSER_VERSION_SPREADSHEET,
        )


class ImageOCRParser:
    supported_file_types = IMAGE_FILE_TYPES

    def __init__(self, backend: OCRParser) -> None:
        self.backend = backend

    def parse_images(self, file: BinaryIO, language_hints: list[str]) -> ParsedDocument:
        document = self.backend.parse_images(file, language_hints)
        # 后端已产出 blocks 则原样透传；否则由共享派生逻辑补齐（保证 OCR 版
        # 面块稳定：bbox / confidence）。
        blocks = document.blocks or _build_ocr_blocks(document)
        return replace(
            document,
            parser_version=PARSER_VERSION_OCR,
            blocks=blocks,
        )
